import os
import sqlite3
import datetime
import urllib.parse
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "store.db")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
SEED_XLSX_PATH = "/home/user/Inventario_Tienda_WhatsApp.xlsx"

os.makedirs(UPLOADS_DIR, exist_ok=True)


def get_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn


def compute_disponibilidad(cantidad: int, estado: str) -> str:
    if estado == "Oculto":
        return "Oculto"
    if estado == "Apartado":
        return "Apartado"
    if cantidad <= 0 or estado == "Vendido":
        return "Agotado"
    return "Disponible"


def build_whatsapp_product_link(whatsapp_num: str, prod: Dict[str, Any]) -> str:
    clean_num = "".join(ch for ch in str(whatsapp_num or "") if ch.isdigit())
    talla_part = f" - Talla: {prod['talla']}" if prod.get("talla") and prod["talla"] not in ("N/A", "Única", "Estándar") else ""
    msg = (
        f"Hola, me interesa comprar {prod['nombre']} "
        f"(ID: {prod['id']}{talla_part} - ${float(prod['precio']):,.0f} MXN)"
    )
    return f"https://wa.me/{clean_num}?text={urllib.parse.quote(msg)}"


def generate_sample_receipt_image(filename: str, folio: str, cliente: str, producto: str, monto: float, metodo: str, fecha: str, estado: str):
    filepath = os.path.join(UPLOADS_DIR, filename)
    if os.path.exists(filepath):
        return f"/uploads/{filename}"

    w, h = 640, 820
    img = Image.new("RGB", (w, h), "#F1F5F9")
    draw = ImageDraw.Draw(img)

    # Card container
    draw.rounded_rectangle((40, 40, w - 40, h - 40), radius=24, fill="#FFFFFF", outline="#CBD5E1", width=2)

    # Top header band
    header_color = "#075E54" if "SPEI" in metodo else "#B45309"
    draw.rounded_rectangle((40, 40, w - 40, 185), radius=24, fill=header_color)
    draw.rectangle((40, 130, w - 40, 185), fill=header_color)

    ttf_reg = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    ttf_bold = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    if os.path.exists(ttf_reg) and os.path.exists(ttf_bold):
        f_title = ImageFont.truetype(ttf_bold, 18)
        f_sub = ImageFont.truetype(ttf_reg, 14)
        f_amt = ImageFont.truetype(ttf_bold, 32)
        f_lbl = ImageFont.truetype(ttf_reg, 14)
        f_val = ImageFont.truetype(ttf_bold, 14)
        f_badge = ImageFont.truetype(ttf_bold, 14)
    else:
        f_title = f_sub = f_amt = f_lbl = f_val = f_badge = ImageFont.load_default()

    draw.text((75, 68), f"COMPROBANTE DE PAGO · {metodo.upper()}", fill="#FFFFFF", font=f_title)
    draw.text((75, 104), f"Referencia de Orden: {folio}", fill="#D1FAE5", font=f_sub)
    draw.text((75, 134), f"Fecha y hora: {fecha} 14:32 hrs", fill="#ECFDF5", font=f_sub)

    # Amount highlight box
    draw.rounded_rectangle((75, 210, w - 75, 328), radius=16, fill="#F8FAFC", outline="#E2E8F0", width=2)
    draw.text((95, 228), "MONTO TRANSFERIDO (MXN)", fill="#64748B", font=f_lbl)
    draw.text((95, 262), f"$ {monto:,.2f} MXN", fill="#0F172A", font=f_amt)

    rows = [
        ("Ordenante / Cliente:", cliente),
        ("Concepto de Pago:", f"{folio} - {producto[:30]}"),
        ("Método Utilizado:", metodo),
        ("Cuenta Destino:", "CLABE **** **** **** 7890"),
        ("Clave de Rastreo:", f"MX2026SPEI{abs(hash(folio)) % 900000 + 100000}"),
        ("Estatus Bancario:", "Operación Exitosa / Enviada"),
    ]

    y = 358
    for label, val in rows:
        draw.text((75, y), label, fill="#64748B", font=f_lbl)
        draw.text((255, y), str(val), fill="#0F172A", font=f_val)
        draw.line((75, y + 34, w - 75, y + 34), fill="#F1F5F9", width=2)
        y += 52

    badge_bg = "#DCFCE7" if estado == "Aprobado" else "#FEF3C7"
    badge_fg = "#166534" if estado == "Aprobado" else "#92400E"
    draw.rounded_rectangle((75, 688, w - 75, 744), radius=12, fill=badge_bg)
    status_txt = "COMPROBANTE VERIFICADO" if estado == "Aprobado" else "PENDIENTE DE VALIDACIÓN"
    draw.text((95, 706), f"Estado en Tienda: {status_txt}", fill=badge_fg, font=f_badge)

    img.save(filepath, "PNG")
    return f"/uploads/{filename}"


def init_db():
    conn = get_conn()
    cur = conn.cursor()

    cur.executescript(
        """
        CREATE TABLE IF NOT EXISTS config (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS categories (
            id TEXT PRIMARY KEY,
            nombre TEXT UNIQUE NOT NULL,
            imagen_url TEXT DEFAULT '',
            activa TEXT DEFAULT 'Sí',
            orden INTEGER DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS products (
            id TEXT PRIMARY KEY,
            nombre TEXT NOT NULL,
            categoria TEXT NOT NULL,
            precio REAL NOT NULL DEFAULT 0,
            cantidad INTEGER NOT NULL DEFAULT 0,
            talla TEXT DEFAULT '',
            color TEXT DEFAULT '',
            marca TEXT DEFAULT '',
            descripcion TEXT DEFAULT '',
            foto_url TEXT DEFAULT '',
            video_url TEXT DEFAULT '',
            condicion TEXT DEFAULT 'Nuevo',
            estado TEXT DEFAULT 'Disponible',
            fecha_alta TEXT NOT NULL,
            destacado TEXT DEFAULT 'No',
            publicado_wa TEXT DEFAULT 'No'
        );

        CREATE TABLE IF NOT EXISTS customers (
            id TEXT PRIMARY KEY,
            nombre TEXT NOT NULL,
            apellidos TEXT DEFAULT '',
            telefono TEXT NOT NULL,
            whatsapp TEXT NOT NULL,
            email TEXT DEFAULT '',
            categorias_interes TEXT DEFAULT '',
            fecha_registro TEXT NOT NULL,
            ultimo_pedido TEXT DEFAULT '',
            estado TEXT DEFAULT 'Activo'
        );

        CREATE TABLE IF NOT EXISTS orders (
            id TEXT PRIMARY KEY,
            fecha TEXT NOT NULL,
            cliente_id TEXT DEFAULT '',
            cliente_nombre TEXT NOT NULL,
            telefono TEXT NOT NULL,
            producto_id TEXT DEFAULT '',
            producto_nombre TEXT NOT NULL,
            cantidad INTEGER NOT NULL DEFAULT 1,
            precio_unitario REAL NOT NULL DEFAULT 0,
            total REAL NOT NULL DEFAULT 0,
            metodo_pago TEXT DEFAULT 'Transferencia SPEI',
            comprobante_url TEXT DEFAULT '',
            estado_pago TEXT DEFAULT 'Pendiente',
            estado_pedido TEXT DEFAULT 'Apartado',
            notas TEXT DEFAULT '',
            stock_deducted INTEGER DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS movements (
            id TEXT PRIMARY KEY,
            fecha TEXT NOT NULL,
            producto_id TEXT NOT NULL,
            producto_nombre TEXT NOT NULL,
            tipo TEXT NOT NULL,
            cambio INTEGER NOT NULL,
            cantidad_anterior INTEGER NOT NULL,
            nueva_cantidad INTEGER NOT NULL,
            motivo TEXT DEFAULT ''
        );

        CREATE TABLE IF NOT EXISTS users (
            id TEXT PRIMARY KEY,
            nombre TEXT NOT NULL,
            email TEXT UNIQUE NOT NULL,
            rol TEXT NOT NULL DEFAULT 'Admin',
            whatsapp TEXT DEFAULT '',
            activo TEXT DEFAULT 'Sí',
            permisos TEXT DEFAULT ''
        );
        """
    )
    conn.commit()

    # Check if already seeded
    count_prod = cur.execute("SELECT COUNT(*) FROM products").fetchone()[0]
    if count_prod == 0:
        seed_from_xlsx_or_defaults(conn)

    conn.close()


def _fmt_cell_date(val) -> str:
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    if val is None or str(val).strip() == "":
        return ""
    s = str(val).strip()
    # Try dd/mm/yyyy
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return datetime.date.today().strftime("%Y-%m-%d")


def seed_from_xlsx_or_defaults(conn: sqlite3.Connection):
    cur = conn.cursor()
    default_cfg = {
        "store_name": "Mi Tienda WhatsApp",
        "whatsapp_number": "525512345678",
        "glide_url": "https://mitienda.glide.page",
        "bank_name": "BBVA México",
        "bank_holder": "Mi Tienda Retail S.A.",
        "bank_clabe": "012180001234567890",
        "oxxo_card": "4152 3138 9900 1234",
        "pickup_info": "Entregas personales en San Mateo Atenco / Toluca y envíos a todo México por paquetería.",
        "admin_pin": "1234",
        "hide_sold_out_default": "false",
    }

    if os.path.exists(SEED_XLSX_PATH):
        wb = openpyxl.load_workbook(SEED_XLSX_PATH, data_only=True)

        if "Dashboard" in wb.sheetnames:
            ws_d = wb["Dashboard"]
            if ws_d["C5"].value:
                default_cfg["store_name"] = str(ws_d["C5"].value).strip()
            if ws_d["C6"].value:
                default_cfg["whatsapp_number"] = str(ws_d["C6"].value).strip()
            if ws_d["C7"].value:
                default_cfg["glide_url"] = str(ws_d["C7"].value).strip()

        for k, v in default_cfg.items():
            cur.execute("INSERT OR REPLACE INTO config (key, value) VALUES (?, ?)", (k, v))

        if "Categorías" in wb.sheetnames:
            ws_c = wb["Categorías"]
            for idx, r in enumerate(range(2, ws_c.max_row + 1), start=1):
                cid = ws_c.cell(r, 1).value
                cname = ws_c.cell(r, 2).value
                if not cid or not cname:
                    continue
                cur.execute(
                    "INSERT OR IGNORE INTO categories (id, nombre, imagen_url, activa, orden) VALUES (?, ?, ?, ?, ?)",
                    (
                        str(cid).strip(),
                        str(cname).strip(),
                        str(ws_c.cell(r, 3).value or "").strip(),
                        str(ws_c.cell(r, 4).value or "Sí").strip(),
                        idx,
                    ),
                )

        if "Productos" in wb.sheetnames:
            ws_p = wb["Productos"]
            for r in range(2, ws_p.max_row + 1):
                pid = ws_p.cell(r, 1).value
                pname = ws_p.cell(r, 2).value
                if not pid or not pname:
                    continue
                cur.execute(
                    """
                    INSERT OR IGNORE INTO products
                    (id, nombre, categoria, precio, cantidad, talla, color, marca, descripcion,
                     foto_url, video_url, condicion, estado, fecha_alta, destacado, publicado_wa)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(pid).strip(),
                        str(pname).strip(),
                        str(ws_p.cell(r, 3).value or "Otros").strip(),
                        float(ws_p.cell(r, 4).value or 0),
                        int(float(ws_p.cell(r, 5).value or 0)),
                        str(ws_p.cell(r, 6).value or "").strip(),
                        str(ws_p.cell(r, 7).value or "").strip(),
                        str(ws_p.cell(r, 8).value or "").strip(),
                        str(ws_p.cell(r, 9).value or "").strip(),
                        str(ws_p.cell(r, 10).value or "").strip(),
                        str(ws_p.cell(r, 11).value or "").strip(),
                        str(ws_p.cell(r, 12).value or "Nuevo").strip(),
                        str(ws_p.cell(r, 13).value or "Disponible").strip(),
                        _fmt_cell_date(ws_p.cell(r, 15).value) or "2026-10-30",
                        str(ws_p.cell(r, 16).value or "No").strip(),
                        str(ws_p.cell(r, 17).value or "No").strip(),
                    ),
                )

        if "Clientes" in wb.sheetnames:
            ws_cl = wb["Clientes"]
            for r in range(2, ws_cl.max_row + 1):
                clid = ws_cl.cell(r, 1).value
                clnom = ws_cl.cell(r, 2).value
                if not clid or not clnom:
                    continue
                cur.execute(
                    """
                    INSERT OR IGNORE INTO customers
                    (id, nombre, apellidos, telefono, whatsapp, email, categorias_interes, fecha_registro, ultimo_pedido, estado)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(clid).strip(),
                        str(clnom).strip(),
                        str(ws_cl.cell(r, 3).value or "").strip(),
                        str(ws_cl.cell(r, 5).value or "").strip(),
                        str(ws_cl.cell(r, 6).value or "").strip(),
                        str(ws_cl.cell(r, 8).value or "").strip(),
                        str(ws_cl.cell(r, 9).value or "").strip(),
                        _fmt_cell_date(ws_cl.cell(r, 10).value) or "2026-10-30",
                        _fmt_cell_date(ws_cl.cell(r, 11).value),
                        str(ws_cl.cell(r, 14).value or "Activo").strip(),
                    ),
                )

        if "Pedidos" in wb.sheetnames:
            ws_o = wb["Pedidos"]
            for r in range(2, ws_o.max_row + 1):
                oid = ws_o.cell(r, 1).value
                if not oid:
                    continue
                oid_str = str(oid).strip()
                fecha_str = _fmt_cell_date(ws_o.cell(r, 2).value) or "2026-10-30"
                cnom = str(ws_o.cell(r, 4).value or "").strip()
                pnom = str(ws_o.cell(r, 7).value or "").strip()
                cant = int(float(ws_o.cell(r, 8).value or 1))
                punit = float(ws_o.cell(r, 9).value or 0)
                tot = cant * punit
                metodo = str(ws_o.cell(r, 11).value or "Transferencia SPEI").strip()
                est_pago = str(ws_o.cell(r, 13).value or "Pendiente").strip()
                est_ped = str(ws_o.cell(r, 14).value or "Apartado").strip()

                # Generate realistic sample receipt PNG so Admin can inspect receipts right away!
                sample_file = f"comprobante_{oid_str.lower().replace('-', '')}.png"
                receipt_url = generate_sample_receipt_image(
                    sample_file, oid_str, cnom, pnom, tot, metodo, fecha_str, est_pago
                )

                # For PED-003 (P004 which has cantidad=1 & estado=Apartado in sheet), stock_deducted=0 so approving deducts 1->0 cleanly
                pid_str = str(ws_o.cell(r, 6).value or "").strip()
                stock_ded = 0 if (pid_str == "P004" and est_pago == "Pendiente") else 1

                cur.execute(
                    """
                    INSERT OR IGNORE INTO orders
                    (id, fecha, cliente_id, cliente_nombre, telefono, producto_id, producto_nombre,
                     cantidad, precio_unitario, total, metodo_pago, comprobante_url, estado_pago, estado_pedido, notas, stock_deducted)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        oid_str,
                        fecha_str,
                        str(ws_o.cell(r, 3).value or "").strip(),
                        cnom,
                        str(ws_o.cell(r, 5).value or "").strip(),
                        pid_str,
                        pnom,
                        cant,
                        punit,
                        tot,
                        metodo,
                        receipt_url,
                        est_pago,
                        est_ped,
                        str(ws_o.cell(r, 15).value or "").strip(),
                        stock_ded,
                    ),
                )

        if "Inventario" in wb.sheetnames:
            ws_i = wb["Inventario"]
            for r in range(2, ws_i.max_row + 1):
                mid = ws_i.cell(r, 1).value
                if not mid:
                    continue
                cur.execute(
                    """
                    INSERT OR IGNORE INTO movements
                    (id, fecha, producto_id, producto_nombre, tipo, cambio, cantidad_anterior, nueva_cantidad, motivo)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(mid).strip(),
                        _fmt_cell_date(ws_i.cell(r, 2).value) or "2026-10-30",
                        str(ws_i.cell(r, 3).value or "").strip(),
                        str(ws_i.cell(r, 4).value or "").strip(),
                        str(ws_i.cell(r, 5).value or "Ajuste de Inventario").strip(),
                        int(float(ws_i.cell(r, 6).value or 0)),
                        int(float(ws_i.cell(r, 7).value or 0)),
                        int(float(ws_i.cell(r, 8).value or 0)),
                        str(ws_i.cell(r, 9).value or "").strip(),
                    ),
                )

        if "Usuarios" in wb.sheetnames:
            ws_u = wb["Usuarios"]
            for r in range(2, ws_u.max_row + 1):
                uid = ws_u.cell(r, 1).value
                uemail = ws_u.cell(r, 3).value
                if not uid or not uemail:
                    continue
                cur.execute(
                    """
                    INSERT OR IGNORE INTO users
                    (id, nombre, email, rol, whatsapp, activo, permisos)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(uid).strip(),
                        str(ws_u.cell(r, 2).value or "").strip(),
                        str(uemail).strip(),
                        str(ws_u.cell(r, 4).value or "Admin").strip(),
                        str(ws_u.cell(r, 5).value or "").strip(),
                        str(ws_u.cell(r, 6).value or "Sí").strip(),
                        str(ws_u.cell(r, 7).value or "").strip(),
                    ),
                )
    else:
        for k, v in default_cfg.items():
            cur.execute("INSERT OR REPLACE INTO config (key, value) VALUES (?, ?)", (k, v))

    conn.commit()


def next_id(conn: sqlite3.Connection, table: str, prefix: str, pad: int = 3) -> str:
    rows = conn.execute(f"SELECT id FROM {table}").fetchall()
    max_n = 0
    for r in rows:
        raw = str(r["id"])
        digits = "".join(ch for ch in raw if ch.isdigit())
        if digits:
            max_n = max(max_n, int(digits))
    return f"{prefix}{str(max_n + 1).zfill(pad)}"


def export_db_to_xlsx(target_path: str) -> str:
    """Exports the current SQLite store state into the exact 7-tab Google Sheet / Excel structure."""
    conn = get_conn()
    cfg = {r["key"]: r["value"] for r in conn.execute("SELECT key, value FROM config").fetchall()}
    products = [dict(r) for r in conn.execute("SELECT * FROM products ORDER BY id ASC").fetchall()]
    customers = [dict(r) for r in conn.execute("SELECT * FROM customers ORDER BY id ASC").fetchall()]
    orders = [dict(r) for r in conn.execute("SELECT * FROM orders ORDER BY id ASC").fetchall()]
    categories = [dict(r) for r in conn.execute("SELECT * FROM categories ORDER BY orden ASC, id ASC").fetchall()]
    movements = [dict(r) for r in conn.execute("SELECT * FROM movements ORDER BY id ASC").fetchall()]
    users = [dict(r) for r in conn.execute("SELECT * FROM users ORDER BY id ASC").fetchall()]
    conn.close()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_prod = wb.create_sheet("Productos")
    ws_cli = wb.create_sheet("Clientes")
    ws_ped = wb.create_sheet("Pedidos")
    ws_cat = wb.create_sheet("Categorías")
    ws_inv = wb.create_sheet("Inventario")
    ws_dash = wb.create_sheet("Dashboard")
    ws_usr = wb.create_sheet("Usuarios")

    ws_prod.sheet_properties.tabColor = "075E54"
    ws_cli.sheet_properties.tabColor = "0284C7"
    ws_ped.sheet_properties.tabColor = "D97706"
    ws_cat.sheet_properties.tabColor = "7C3AED"
    ws_inv.sheet_properties.tabColor = "475569"
    ws_dash.sheet_properties.tabColor = "10B981"
    ws_usr.sheet_properties.tabColor = "DC2626"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color="0F172A")
    bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    fill_hdr = PatternFill(start_color="075E54", end_color="075E54", fill_type="solid")
    fill_auto = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    thin_side = Side(border_style="thin", color="CBD5E1")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    def setup_headers(ws, headers, auto_cols, widths):
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 28
        for idx, h in enumerate(headers, 1):
            c = ws.cell(1, idx, h)
            c.font = header_font
            c.fill = fill_auto if idx in auto_cols else fill_hdr
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]

    # 1. Productos
    setup_headers(
        ws_prod,
        ["ID Producto", "Nombre del Producto", "Categoría", "Precio (MXN)", "Cantidad", "Talla / Medida",
         "Color", "Marca", "Descripción", "URL Foto", "URL Video", "Condición", "Estado",
         "Disponibilidad Automática", "Fecha de Alta", "Destacado", "Publicado en WhatsApp", "Enlace WhatsApp"],
        {14, 18},
        [14, 30, 22, 15, 11, 16, 15, 16, 36, 28, 26, 14, 15, 22, 15, 13, 20, 42]
    )
    for idx, p in enumerate(products, start=2):
        row = [
            p["id"], p["nombre"], p["categoria"], p["precio"], p["cantidad"], p["talla"],
            p["color"], p["marca"], p["descripcion"], p["foto_url"], p["video_url"],
            p["condicion"], p["estado"],
            f'=IF(A{idx}="","",IF(M{idx}="Oculto","Oculto",IF(M{idx}="Apartado","Apartado",IF(OR(E{idx}<=0,M{idx}="Vendido"),"Agotado","Disponible"))))',
            p["fecha_alta"], p["destacado"], p["publicado_wa"],
            f'=IF(B{idx}="","","https://wa.me/"&Dashboard!$C$6&"?text=Hola%2C%20me%20interesa%20comprar%20"&SUBSTITUTE(TRIM(B{idx})," ","%20")&"%20(ID%3A%20"&A{idx}&"%20-%20%24"&TEXT(D{idx},"#,##0")&"%20MXN)")'
        ]
        for c_i, val in enumerate(row, 1):
            cell = ws_prod.cell(idx, c_i, val)
            cell.font = body_font
            cell.border = border
            if c_i == 4:
                cell.number_format = "$#,##0.00"

    # 2. Clientes
    setup_headers(
        ws_cli,
        ["ID Cliente", "Nombre(s)", "Apellidos", "Nombre Completo", "Teléfono", "WhatsApp",
         "Chat Directo WhatsApp", "Correo Electrónico", "Categorías de Interés", "Fecha de Registro",
         "Último Pedido", "Total de Pedidos", "Total Comprado (MXN)", "Estado"],
        {4, 7, 12, 13},
        [13, 16, 20, 24, 16, 16, 30, 26, 32, 17, 16, 17, 20, 14]
    )
    for idx, c in enumerate(customers, start=2):
        row = [
            c["id"], c["nombre"], c["apellidos"],
            f'=IF(AND(B{idx}="",C{idx}=""),"",TRIM(B{idx}&" "&C{idx}))',
            c["telefono"], c["whatsapp"],
            f'=IF(F{idx}="","","https://wa.me/"&SUBSTITUTE(F{idx}," ",""))',
            c["email"], c["categorias_interes"], c["fecha_registro"], c["ultimo_pedido"],
            f'=IF(A{idx}="","",COUNTIF(Pedidos!$C$2:$C$500,A{idx}))',
            f'=IF(A{idx}="","",SUMIFS(Pedidos!$J$2:$J$500,Pedidos!$C$2:$C$500,A{idx},Pedidos!$M$2:$M$500,"Aprobado"))',
            c["estado"]
        ]
        for c_i, val in enumerate(row, 1):
            cell = ws_cli.cell(idx, c_i, val)
            cell.font = body_font
            cell.border = border
            if c_i == 13:
                cell.number_format = "$#,##0.00"

    # 3. Pedidos
    setup_headers(
        ws_ped,
        ["ID Pedido", "Fecha", "ID Cliente", "Nombre del Cliente", "Teléfono", "ID Producto",
         "Nombre del Producto", "Cantidad", "Precio Unitario (MXN)", "Total del Pedido (MXN)",
         "Método de Pago", "URL Comprobante", "Estado de Pago", "Estado del Pedido", "Notas"],
        {10},
        [13, 14, 13, 24, 16, 13, 28, 11, 20, 21, 24, 30, 16, 20, 34]
    )
    for idx, o in enumerate(orders, start=2):
        row = [
            o["id"], o["fecha"], o["cliente_id"], o["cliente_nombre"], o["telefono"],
            o["producto_id"], o["producto_nombre"], o["cantidad"], o["precio_unitario"],
            f'=IF(OR(H{idx}="",I{idx}=""),"",H{idx}*I{idx})',
            o["metodo_pago"], o["comprobante_url"], o["estado_pago"], o["estado_pedido"], o["notas"]
        ]
        for c_i, val in enumerate(row, 1):
            cell = ws_ped.cell(idx, c_i, val)
            cell.font = body_font
            cell.border = border
            if c_i in (9, 10):
                cell.number_format = "$#,##0.00"

    # 4. Categorías
    setup_headers(
        ws_cat,
        ["ID Categoría", "Nombre de Categoría", "URL Imagen", "Activa", "Productos Registrados", "Unidades Disponibles"],
        {5, 6},
        [15, 24, 44, 12, 22, 22]
    )
    for idx, cat in enumerate(categories, start=2):
        row = [
            cat["id"], cat["nombre"], cat["imagen_url"], cat["activa"],
            f'=IF(B{idx}="","",COUNTIF(Productos!$C$2:$C$500,B{idx}))',
            f'=IF(B{idx}="","",SUMIFS(Productos!$E$2:$E$500,Productos!$C$2:$C$500,B{idx},Productos!$N$2:$N$500,"Disponible"))'
        ]
        for c_i, val in enumerate(row, 1):
            cell = ws_cat.cell(idx, c_i, val)
            cell.font = body_font
            cell.border = border

    # 5. Inventario
    setup_headers(
        ws_inv,
        ["ID Movimiento", "Fecha", "ID Producto", "Nombre del Producto", "Tipo de Movimiento",
         "Cambio de Cantidad", "Cantidad Anterior", "Nueva Cantidad", "Motivo / Referencia"],
        {4, 8},
        [16, 14, 14, 30, 24, 19, 18, 16, 36]
    )
    for idx, m in enumerate(movements, start=2):
        row = [
            m["id"], m["fecha"], m["producto_id"],
            f'=IF(C{idx}="","",IFERROR(VLOOKUP(C{idx},Productos!$A$2:$B$500,2,FALSE),"{m["producto_nombre"]}"))',
            m["tipo"], m["cambio"], m["cantidad_anterior"],
            f'=IF(AND(F{idx}="",G{idx}=""),"",G{idx}+F{idx})',
            m["motivo"]
        ]
        for c_i, val in enumerate(row, 1):
            cell = ws_inv.cell(idx, c_i, val)
            cell.font = body_font
            cell.border = border

    # 6. Dashboard
    ws_dash.column_dimensions["B"].width = 38
    ws_dash.column_dimensions["C"].width = 24
    ws_dash["B2"] = "PANEL DE CONTROL — TIENDA WHATSAPP"
    ws_dash["B2"].font = Font(name="Calibri", size=14, bold=True, color="075E54")
    ws_dash["B5"] = "Nombre de la Tienda"; ws_dash["C5"] = cfg.get("store_name", "Mi Tienda WhatsApp")
    ws_dash["B6"] = "Número de WhatsApp (con clave país 52)"; ws_dash["C6"] = cfg.get("whatsapp_number", "525512345678")
    ws_dash["C6"].number_format = "@"
    ws_dash["B7"] = "Enlace Público de tu App"; ws_dash["C7"] = cfg.get("glide_url", "")

    kpis = [
        (11, "Total de Productos Registrados", "=COUNTA(Productos!A2:A500)", "#,##0"),
        (12, "Productos Disponibles", '=COUNTIF(Productos!N2:N500,"Disponible")', "#,##0"),
        (13, "Productos Apartados", '=COUNTIF(Productos!N2:N500,"Apartado")', "#,##0"),
        (14, "Productos Vendidos / Agotados", '=COUNTIF(Productos!N2:N500,"Agotado")', "#,##0"),
        (15, "Unidades Totales Disponibles", '=SUMIF(Productos!N2:N500,"Disponible",Productos!E2:E500)', "#,##0"),
        (16, "Valor del Inventario Disponible (MXN)", '=SUMPRODUCT(--(Productos!N2:N500="Disponible"),Productos!D2:D500,Productos!E2:E500)', "$#,##0.00"),
        (17, "Pendientes de Publicar en Estado WhatsApp", '=COUNTIFS(Productos!Q2:Q500,"No",Productos!N2:N500,"Disponible")', "#,##0"),
        (18, "Pagos Pendientes de Revisión", '=COUNTIF(Pedidos!M2:M500,"Pendiente")', "#,##0"),
        (19, "Monto en Pagos Pendientes (MXN)", '=SUMIFS(Pedidos!J2:J500,Pedidos!M2:M500,"Pendiente")', "$#,##0.00"),
        (20, "Ventas Cobradas / Aprobadas (MXN)", '=SUMIFS(Pedidos!J2:J500,Pedidos!M2:M500,"Aprobado")', "$#,##0.00"),
        (21, "Total de Pedidos Registrados", "=COUNTA(Pedidos!A2:A500)", "#,##0"),
        (22, "Total de Clientes en Base de Datos", "=COUNTA(Clientes!A2:A500)", "#,##0"),
    ]
    for r_i, lbl, form, fmt in kpis:
        c1 = ws_dash.cell(r_i, 2, lbl); c1.font = body_font; c1.border = border
        c2 = ws_dash.cell(r_i, 3, form); c2.font = bold_font; c2.border = border; c2.number_format = fmt

    # 7. Usuarios
    setup_headers(
        ws_usr,
        ["ID Usuario", "Nombre", "Correo Electrónico", "Rol", "WhatsApp", "Activo", "Permisos en Glide"],
        set(),
        [14, 26, 28, 15, 18, 12, 52]
    )
    for idx, u in enumerate(users, start=2):
        row = [u["id"], u["nombre"], u["email"], u["rol"], u["whatsapp"], u["activo"], u["permisos"]]
        for c_i, val in enumerate(row, 1):
            cell = ws_usr.cell(idx, c_i, val)
            cell.font = body_font
            cell.border = border

    wb.save(target_path)
    return target_path
